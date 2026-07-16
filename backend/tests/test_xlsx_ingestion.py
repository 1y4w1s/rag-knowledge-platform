"""F1：Excel .xlsx 入库管道 & 表格 chunk 测试。"""

from __future__ import annotations

import uuid
from pathlib import Path

import pytest
from httpx import AsyncClient
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.enums import DocumentStatus
from app.services.ingestion.pipeline import process_document_ingestion

from tests.conftest import create_test_kb as _create_kb

pytestmark = pytest.mark.asyncio


def _make_golden_xlsx(path: Path) -> None:
    """Create a small test xlsx with 2 sheets containing table data."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()

    # Sheet 1: 部门预算
    ws1 = wb.active
    ws1.title = "部门预算"
    ws1.append(["部门", "预算(万)", "实际(万)"])
    ws1.append(["研发部", 500, 480])
    ws1.append(["市场部", 300, 310])
    ws1.append(["行政部", 100, 95])

    # Sheet 2: 项目支出
    ws2 = wb.create_sheet("项目支出")
    ws2.append(["项目", "负责人", "金额(万)", "状态"])
    ws2.append(["睿阁 v1", "张三", 200, "已完成"])
    ws2.append(["OCR 模块", "李四", 80, "进行中"])
    ws2.append(["数据迁移", "王五", 50, "未开始"])

    wb.save(str(path))


async def _load_chunks(document_id: uuid.UUID) -> list[DocumentChunk]:
    async with SessionLocal() as db:
        result = await db.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
            .order_by(DocumentChunk.chunk_index)
        )
        return list(result.all())


@pytest.fixture
def upload_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    from app.core.config import settings
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    return tmp_path


async def _ingest_fixture(
    *,
    kb_id: uuid.UUID,
    user_id: uuid.UUID,
    source: Path,
    file_type: str,
    upload_dir: Path,
) -> Document:
    doc_id = uuid.uuid4()
    storage_dir = upload_dir / str(kb_id) / str(doc_id)
    storage_dir.mkdir(parents=True, exist_ok=True)
    storage_path = storage_dir / f"{uuid.uuid4()}.{file_type}"
    storage_path.write_bytes(source.read_bytes())

    async with SessionLocal() as db:
        doc = Document(
            id=doc_id,
            kb_id=kb_id,
            filename=source.name,
            file_type=file_type,
            file_size=storage_path.stat().st_size,
            storage_path=str(storage_path),
            status=DocumentStatus.queued,
            uploaded_by=user_id,
        )
        db.add(doc)
        await db.commit()

    await process_document_ingestion(doc_id)

    async with SessionLocal() as db:
        row = await db.get(Document, doc_id)
        assert row is not None
        return row


async def test_xlsx_ingestion_creates_table_chunks(
    client: AsyncClient,
    register_and_login,
    upload_dir: Path,
    tmp_path: Path,
) -> None:
    """xlsx 文件入库后生成 table chunk，每 sheet 一张表。"""
    pytest.importorskip("openpyxl")
    xlsx_path = tmp_path / "golden_ledger.xlsx"
    _make_golden_xlsx(xlsx_path)

    headers, user = await register_and_login(prefix="xlsx-test")
    kb = await _create_kb(client, headers, user)
    kb_id = uuid.UUID(kb["id"])
    user_id = uuid.UUID(user["id"])

    doc = await _ingest_fixture(
        kb_id=kb_id,
        user_id=user_id,
        source=xlsx_path,
        file_type="xlsx",
        upload_dir=upload_dir,
    )

    assert doc.status == DocumentStatus.completed
    assert doc.chunk_count is not None and doc.chunk_count > 0
    assert doc.processing_completed_at is not None

    chunks = await _load_chunks(doc.id)
    assert len(chunks) == doc.chunk_count
    assert all(c.embedding is not None for c in chunks)

    # Should have 2 table chunks (one per sheet)
    table_chunks = [c for c in chunks if c.chunk_kind == "table"]
    assert len(table_chunks) >= 2, f"Expected >=2 table chunks, got {len(table_chunks)}"

    # Verify sheet names appear in heading_path
    heading_paths = {c.heading_path for c in table_chunks}
    assert "部门预算" in heading_paths, f"Missing '部门预算' in {heading_paths}"
    assert "项目支出" in heading_paths, f"Missing '项目支出' in {heading_paths}"

    # Verify content has table data
    contents = " ".join(c.content for c in table_chunks)
    assert "研发部" in contents
    assert "睿阁 v1" in contents or "睿阁" in contents


async def test_xlsx_ingestion_upload_endpoint(
    client: AsyncClient,
    register_and_login,
    tmp_path: Path,
) -> None:
    """通过上传 API 验证 xlsx 可正常上传入库。"""
    pytest.importorskip("openpyxl")
    xlsx_path = tmp_path / "test_upload.xlsx"
    _make_golden_xlsx(xlsx_path)

    headers, user = await register_and_login(prefix="xlsx-upload")
    kb = await _create_kb(client, headers, user)
    kb_id = kb["id"]

    # Upload via API
    with open(xlsx_path, "rb") as f:
        upload_resp = await client.post(
            f"/api/v1/knowledge-bases/{kb_id}/documents",
            files={"files": ("test_ledger.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
    assert upload_resp.status_code == 201, upload_resp.text
    data = upload_resp.json()
    docs = data["documents"]
    assert len(docs) == 1
    assert docs[0]["file_type"] == "xlsx"
    assert docs[0]["status"] in ("queued", "completed")
